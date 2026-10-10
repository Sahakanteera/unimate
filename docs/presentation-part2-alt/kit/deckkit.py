"""Tiny layout kit: write editable PPTX slides in 1920x1080 "px" coordinates.

Every element is a native PowerPoint object (text box, rectangle, ellipse,
straight connector with arrowhead, picture) so PowerPoint and Canva can edit
each piece separately. Thai text gets its font in both the Latin and the
complex-script (cs) slot, otherwise PowerPoint falls back to a system Thai font.

A manifest of every text box is written next to the .pptx; check.py uses it
to detect text that renders outside its box.
"""
from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 6350                      # EMU per px: 1920 px = 13.333 in
W, H = 1920, 1080
ALIGN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
ANCHOR = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}


def e(v: float) -> Emu:
    return Emu(int(round(v * PX)))


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.lstrip('#').upper())


def _strip_style(shape) -> None:
    """Drop the theme style python-pptx attaches (it can add shadows and white text)."""
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


class Deck:
    def __init__(self, head_font: str, body_font: str, title: str = '', author: str = ''):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = e(W), e(H)
        self.layout = self.prs.slide_layouts[6]
        self.head_font, self.body_font = head_font, body_font
        self.manifest: list[dict] = []
        cp = self.prs.core_properties
        cp.title, cp.author, cp.language = title, author, 'th-TH'
        self._set_theme_fonts()

    def _set_theme_fonts(self) -> None:
        """Theme heading/body fonts, so text added later in PowerPoint/Canva matches."""
        master = self.prs.slide_masters[0]
        theme_part = next(r.target_part for r in master.part.rels.values() if r.reltype.endswith('/theme'))
        root = etree.fromstring(theme_part.blob)
        a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        for tag, font in (('majorFont', self.head_font), ('minorFont', self.body_font)):
            node = root.find(f'.//{{{a}}}{tag}')
            for slot in ('latin', 'cs'):
                el = node.find(f'{{{a}}}{slot}')
                el.set('typeface', font)
            for f in node.findall(f'{{{a}}}font'):
                if f.get('script') == 'Thai':
                    f.set('typeface', font)
        theme_part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

    def slide(self, bg: str = 'FFFFFF', _register: bool = True) -> 'Slide':
        s = self.prs.slides.add_slide(self.layout)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg)
        sl = Slide(self, s, len(self.prs.slides))
        sl.bg = bg
        if _register:
            self.__dict__.setdefault('originals', []).append(sl)
        return sl

    def duplicate(self, src: 'Slide') -> 'Slide':
        """Copy every shape (pictures keep their image) of src onto a new slide."""
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        dst = self.slide(src.bg, _register=False)
        tree = dst.s.shapes._spTree
        for el in src.s.shapes._spTree.iterchildren():
            if etree.QName(el).localname in ('nvGrpSpPr', 'grpSpPr', 'extLst'):
                continue
            new = deepcopy(el)
            for blip in new.iter(qn('a:blip')):
                part = src.s.part.related_part(blip.get(qn('r:embed')))
                blip.set(qn('r:embed'), dst.s.part.relate_to(part, RT.IMAGE))
            tree.append(new)
        if src.s.has_notes_slide:
            dst.notes(src.s.notes_slide.notes_text_frame.text)
        dst._n = src._n
        return dst

    def _expand_steps(self) -> None:
        """Slides with deferred callouts become: clean slide, then one slide per point.
        Page numbers ('เลขหน้า') and 'สไลด์ N' references are rewritten to the final order."""
        import re
        originals = self.__dict__.get('originals', [])
        if not any(getattr(s, 'deferred', None) for s in originals):
            return
        order, span = [], {}
        for sl in originals:
            group = [sl]
            co = getattr(sl, 'deferred', None)
            if co:
                for k in range(1, len(co.items) + 1):
                    dup = self.duplicate(sl)
                    co.render(dup, k)
                    group.append(dup)
            span[sl.index] = (len(order) + 1, len(order) + len(group))
            order += group
        lst = self.prs.slides._sldIdLst
        by_id = {el.id: el for el in list(lst)}
        ids = [sl.s.slide_id for sl in order]
        for el in list(lst):
            lst.remove(el)
        for sid in ids:
            lst.append(by_id[sid])
        total = len(order)

        def ref(m):
            a = span[int(m.group(1))][0]
            if m.group(3):
                return f'สไลด์ {a}–{span[int(m.group(3))][1]}'
            return f'สไลด์ {a}'
        for i, sl in enumerate(order, 1):
            for sh in sl.s.shapes:
                if not sh.has_text_frame:
                    continue
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if sh.name == 'เลขหน้า':
                            r.text = f'{i:02d} / {total}'
                        elif 'สไลด์ ' in r.text:
                            r.text = re.sub(r'สไลด์ (\d+)(–(\d+))?', ref, r.text)

    def _shape_manifest(self) -> list[dict]:
        out = []
        for i, s in enumerate(self.prs.slides, 1):
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    out.append({'slide': i, 'name': sh.name, 'rect': [sh.left / PX, sh.top / PX, sh.width / PX, sh.height / PX],
                                'text': sh.text_frame.text.strip()})
        return out

    def save(self, path: str | Path) -> None:
        path = Path(path)
        self._expand_steps()
        self.prs.save(path)
        path.with_suffix('.manifest.json').write_text(json.dumps(self._shape_manifest(), ensure_ascii=False, indent=1), encoding='utf-8')


class Slide:
    def __init__(self, deck: Deck, slide, index: int):
        self.deck, self.s, self.index = deck, slide, index
        self._n = 0

    # ---------- helpers ----------
    def _name(self, shape, name: str | None, kind: str) -> None:
        self._n += 1
        shape.name = name or f'{kind} {self._n}'

    def _fill_line(self, shape, fill, line, lw, dash=None, alpha=None) -> None:
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
            if alpha is not None:   # 0..1 opacity
                clr = shape.fill._xPr.find(qn('a:solidFill'))[0]
                etree.SubElement(clr, qn('a:alpha')).set('val', str(int(alpha * 100000)))
        else:
            shape.fill.background()
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = e(lw)
            if dash:
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            shape.line.fill.background()

    # ---------- shapes ----------
    def rect(self, x, y, w, h, fill=None, line=None, lw=2, radius=0, dash=False, alpha=None, name=None):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        sh = self.s.shapes.add_shape(kind, e(x), e(y), e(w), e(h))
        _strip_style(sh)
        if radius:
            sh.adjustments[0] = min(0.5, radius / min(w, h))
        self._fill_line(sh, fill, line, lw, dash, alpha)
        self._name(sh, name, 'Rect')
        return sh

    def ellipse(self, x, y, w, h, fill=None, line=None, lw=2, name=None):
        sh = self.s.shapes.add_shape(MSO_SHAPE.OVAL, e(x), e(y), e(w), e(h))
        _strip_style(sh)
        self._fill_line(sh, fill, line, lw)
        self._name(sh, name, 'Ellipse')
        return sh

    def line(self, x1, y1, x2, y2, color, lw=3, head=None, tail=None, dash=False, name=None):
        """Straight connector. tail = arrowhead at (x2, y2), head = at (x1, y1)."""
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, e(x1), e(y1), e(x2), e(y2))
        _strip_style(c)
        c.line.color.rgb = rgb(color)
        c.line.width = e(lw)
        if dash:
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ln = c.line._get_or_add_ln()
        for tag, kind in (('a:headEnd', head), ('a:tailEnd', tail)):
            if kind:
                el = etree.SubElement(ln, qn(tag))
                el.set('type', kind)
                el.set('w', 'med')
                el.set('len', 'med')
        self._name(c, name, 'Arrow' if (head or tail) else 'Line')
        return c

    def image(self, path, x, y, w, h, line=None, lw=2, name=None):
        pic = self.s.shapes.add_picture(str(path), e(x), e(y), e(w), e(h))
        if line:
            pic.line.color.rgb = rgb(line)
            pic.line.width = e(lw)
        self._name(pic, name, 'Picture')
        return pic

    # ---------- text ----------
    def text(self, x, y, w, h, paras, font=None, size=28, color='000000', bold=False, align='l', anchor='t',
             lh=1.2, after=0, fill=None, line=None, lw=2, radius=0, pad=0, name=None, check=True):
        """paras: str ('\n' = new paragraph) or list of paragraphs; a paragraph is a str
        or a list of runs (text, {font,size,color,bold}). size is px on the 1920 canvas."""
        font = font or self.deck.body_font
        if fill or line:
            sh = self.rect(x, y, w, h, fill=fill, line=line, lw=lw, radius=radius, name=name)
        else:
            sh = self.s.shapes.add_textbox(e(x), e(y), e(w), e(h))
            self._name(sh, name, 'Text')
        tf = sh.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = e(pad if not isinstance(pad, tuple) else pad[0])
        tf.margin_top = tf.margin_bottom = e(pad if not isinstance(pad, tuple) else pad[1])
        tf.vertical_anchor = ANCHOR[anchor]
        if isinstance(paras, str):
            paras = paras.split('\n')
        plain = []
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            p.line_spacing = lh
            if after and i < len(paras) - 1:
                p.space_after = Pt(after * 0.5)
            runs = [(para, {})] if isinstance(para, str) else para
            for txt, st in runs:
                r = p.add_run()
                r.text = txt
                f = r.font
                fn = st.get('font', font)
                f.name = fn
                f.size = Pt(st.get('size', size) * 0.5)
                f.bold = st.get('bold', bold)
                f.color.rgb = rgb(st.get('color', color))
                rpr = r._r.get_or_add_rPr()
                rpr.set('lang', 'th-TH')
                rpr.set('altLang', 'en-US')
                for slot in ('a:ea', 'a:cs'):
                    el = rpr.find(qn(slot))
                    if el is None:
                        el = etree.SubElement(rpr, qn(slot))
                    el.set('typeface', fn)
                plain.append(txt)
            plain.append('\n')
        if check:
            self.deck.manifest.append({'slide': self.index, 'name': sh.name, 'rect': [x, y, w, h], 'text': ''.join(plain).strip()})
        return sh

    def pin(self, n, cx, cy, d=52, fill='E8470C', color='FFFFFF', font=None, size=28, name=None):
        sh = self.text(cx - d / 2, cy - d / 2, d, d, str(n), font=font or self.deck.head_font, size=size, color=color,
                       bold=True, align='c', anchor='m', lh=1.0, fill=fill, pad=0, name=name or f'Pin {n}')
        # turn the rectangle into an ellipse so the number sits in a circle
        sh._element.spPr.find(qn('a:prstGeom')).set('prst', 'ellipse')
        return sh

    def notes(self, text: str) -> None:
        self.s.notes_slide.notes_text_frame.text = text


def text_width(text: str, font_path: str, size_px: float) -> float:
    """Rough rendered width in px (PIL, no Thai shaping: marks counted ~0)."""
    from PIL import ImageFont
    return ImageFont.truetype(font_path, int(size_px)).getlength(text)
